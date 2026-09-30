"""Cierre de dia (fix 2026-09-30): no archivar la fecha en curso,
reparar archivos vacios, Excel consistente y archivado del dia viejo
en el cambio de fecha (seleccionar_partidos._archivar_dia_anterior)."""
import json

import pytest

import cerrar_resultados as C
import seleccionar_partidos as S


def _partido(i, snaps=0, resuelto=False):
    return {
        "fixture_id": None,  # sin fixture: cerrar() no consulta ESPN
        "partido": f"Equipo{i} vs Rival{i}",
        "favorito": f"Equipo{i}",
        "favorito_es_local": True,
        "resultado_final": "2-1" if resuelto else None,
        "acierto": True if resuelto else None,
        "historial_snapshots": [{"minuto": "10'", "goles_local": 0,
                                 "goles_visitante": 0}] * snaps,
        "alertas_enviadas": [],
    }


def _datos(fecha, n=2, snaps=1, resueltos=0):
    return {"fecha": fecha,
            "partidos": [_partido(i, snaps=snaps, resuelto=i < resueltos)
                         for i in range(n)]}


@pytest.fixture
def entorno(tmp_path, monkeypatch):
    monkeypatch.setattr(C, "ARCHIVO_PARTIDOS", tmp_path / "partidos_hoy.json")
    monkeypatch.setattr(C, "DIR_HISTORIAL_DIAS", tmp_path / "historial_dias")
    monkeypatch.setattr(C, "ARCHIVO_EXCEL", tmp_path / "estadisticas.xlsx")
    monkeypatch.setattr(C, "uso_de_hoy", lambda: (7, None))
    marcas = []
    monkeypatch.setattr(C, "marcar_hecho", marcas.append)
    return {"dir": tmp_path, "marcas": marcas}


def test_no_cierra_la_fecha_en_curso(entorno):
    """El cron de Fase 4 atrasado no debe archivar el dia que aun corre."""
    fecha = C._fecha_local_hoy()
    C.cerrar(datos=_datos(fecha))
    assert not (entorno["dir"] / "historial_dias" / f"{fecha}.json").exists()
    assert entorno["marcas"] == []


def test_archiva_un_dia_terminado(entorno):
    C.cerrar(datos=_datos("2000-01-01", n=3, snaps=2, resueltos=1),
             escribir_partidos=False)
    archivo = entorno["dir"] / "historial_dias" / "2000-01-01.json"
    datos = json.loads(archivo.read_text(encoding="utf-8"))
    assert datos["fecha"] == "2000-01-01"
    assert len(datos["partidos"]) == 3
    assert entorno["marcas"] == ["cierre"]
    assert not C.ARCHIVO_PARTIDOS.exists()


def test_sin_partidos_hoy_no_hace_nada(entorno):
    C.cerrar()
    assert entorno["marcas"] == []
    assert not (entorno["dir"] / "historial_dias").exists()


def test_no_reescribe_si_el_archivo_ya_esta_rico(entorno):
    archivo = entorno["dir"] / "historial_dias" / "2000-01-02.json"
    archivo.parent.mkdir(parents=True)
    archivo.write_text(json.dumps(_datos("2000-01-02", n=2, snaps=5)),
                       encoding="utf-8")
    C.cerrar(datos=_datos("2000-01-02", n=2, snaps=1))
    datos = json.loads(archivo.read_text(encoding="utf-8"))
    assert len(datos["partidos"][0]["historial_snapshots"]) == 5
    assert entorno["marcas"] == []
    assert not C.ARCHIVO_EXCEL.exists()


def test_repara_un_archivo_vacio(entorno):
    archivo = entorno["dir"] / "historial_dias" / "2000-01-03.json"
    archivo.parent.mkdir(parents=True)
    archivo.write_text(json.dumps(_datos("2000-01-03", n=2, snaps=0)),
                       encoding="utf-8")
    C.cerrar(datos=_datos("2000-01-03", n=2, snaps=3, resueltos=1))
    datos = json.loads(archivo.read_text(encoding="utf-8"))
    assert sum(len(p["historial_snapshots"]) for p in datos["partidos"]) == 6
    assert sum(1 for p in datos["partidos"] if p["resultado_final"]) == 1
    assert entorno["marcas"] == ["cierre"]


def test_reparacion_no_duplica_filas_en_excel(entorno):
    import openpyxl
    C.cerrar(datos=_datos("2000-01-04", n=2, snaps=1))
    C.cerrar(datos=_datos("2000-01-04", n=2, snaps=4, resueltos=2))
    wb = openpyxl.load_workbook(C.ARCHIVO_EXCEL)
    hoja1 = wb["Resultados diarios"]
    hoja2 = wb["Resumen por dia"]
    filas1 = [r for r in hoja1.iter_rows(min_row=2) if r[0].value == "2000-01-04"]
    filas2 = [r for r in hoja2.iter_rows(min_row=2) if r[0].value == "2000-01-04"]
    assert len(filas1) == 2  # una por partido, sin duplicados
    assert len(filas2) == 1
    assert filas2[0][1].value == 2  # total de partidos actualizado


def test_archivado_en_cambio_de_fecha_no_pisa_partidos_hoy(entorno, monkeypatch):
    """La ruta de cambio de fecha (seleccionar) archiva los datos viejos
    sin escribir de vuelta partidos_hoy.json."""
    partido = _partido(0)
    partido.update({"fixture_id": "123", "liga_slug": "eng.1",
                    "acierto": None, "resultado_final": None})
    datos = {"fecha": "2000-01-05", "partidos": [partido]}
    monkeypatch.setattr(C, "obtener_resultado_fixture",
                        lambda fid, slug: {"fixture": {"status": {"short": "FT"}},
                                           "goals": {"home": 2, "away": 1}})
    monkeypatch.setattr(C, "_actualizar_rating_propio", lambda *a: None)
    C.ARCHIVO_PARTIDOS.write_text(
        json.dumps({"fecha": "2000-01-05", "partidos": []}), encoding="utf-8")

    C.cerrar(datos=datos, escribir_partidos=False)

    en_archivo = json.loads(C.ARCHIVO_PARTIDOS.read_text(encoding="utf-8"))
    assert en_archivo["partidos"] == []  # no se piso el archivo actual
    archivado = json.loads(
        (entorno["dir"] / "historial_dias" / "2000-01-05.json")
        .read_text(encoding="utf-8"))
    assert archivado["partidos"][0]["resultado_final"] == "2-1"


def test_archivar_dia_anterior_solo_si_cambio_la_fecha(monkeypatch):
    llamadas = []

    def fake_cerrar(datos=None, escribir_partidos=True):
        llamadas.append((datos.get("fecha"), escribir_partidos))

    monkeypatch.setattr("cerrar_resultados.cerrar", fake_cerrar)
    hoy = S.fecha_local_hoy()

    S._archivar_dia_anterior({"fecha": "2000-01-01", "partidos": [_partido(0)]})
    assert llamadas == [("2000-01-01", False)]

    llamadas.clear()
    S._archivar_dia_anterior({"fecha": hoy, "partidos": [_partido(0)]})
    S._archivar_dia_anterior({"fecha": "2000-01-01", "partidos": []})
    S._archivar_dia_anterior(None)
    assert llamadas == []
