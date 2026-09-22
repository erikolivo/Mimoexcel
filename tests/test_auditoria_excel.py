"""T6 — Auditoria nocturna y Excel: sin 0 % espurios, alias unificado."""
import copy

import cerrar_resultados as C


def partido_con_alertas():
    return {
        "partido": "Sevilla vs Betis", "favorito_es_local": True,
        "resultado_final": "1-0",
        "historial_snapshots": [
            {"minuto": "34'", "goles_local": 0, "goles_visitante": 0,
             "stats_local": {}, "stats_visitante": {}},
            {"minuto": "41'", "goles_local": 1, "goles_visitante": 0,
             "stats_local": {}, "stats_visitante": {}},
        ],
        "alertas_enviadas": [
            {"tipo": "posible_victoria_favorito", "minuto": "34'", "texto": "x",
             "id": "1-0", "minuto_int": 34, "lado": "fav", "criterio": "siguiente_gol",
             "estado": "pendiente", "acierto": None},
            {"tipo": "no_fav_domina", "minuto": "20'", "texto": "x",
             "id": "1-1", "minuto_int": 20, "lado": "rival", "criterio": "siguiente_gol",
             "estado": "pendiente", "acierto": None},
            {"tipo": "tarjeta_roja", "minuto": "50'", "texto": "x",
             "id": "1-2", "minuto_int": 50, "lado": None, "criterio": None,
             "estado": "no_aplica", "acierto": None},
        ],
    }


def test_auditoria_resuelve_todos_los_tipos_con_criterio():
    p = partido_con_alertas()
    C._auditar_alertas(p)
    por_tipo = {a["tipo"]: a["acierto"] for a in p["alertas_enviadas"]}
    assert por_tipo["posible_victoria_favorito"] is True
    assert por_tipo["no_fav_domina"] is False
    assert por_tipo["tarjeta_roja"] is None


def test_auditoria_respeta_lo_resuelto_en_vivo():
    p = partido_con_alertas()
    p["alertas_enviadas"][0]["estado"] = "fallo"
    p["alertas_enviadas"][0]["acierto"] = False
    C._auditar_alertas(p)
    assert p["alertas_enviadas"][0]["acierto"] is False  # no se recalcula


def test_excel_unifica_alias_y_no_cuenta_none(tmp_path, monkeypatch):
    import openpyxl
    xlsx = tmp_path / "est.xlsx"
    monkeypatch.setattr(C, "ARCHIVO_EXCEL", xlsx)
    p = partido_con_alertas()
    p["acierto"] = True
    p["favorito"] = "Sevilla"
    p["cuota_inicial"] = ""
    p["probabilidad_inicial"] = ""
    p2 = copy.deepcopy(p)
    p2["alertas_enviadas"] = [
        {"tipo": "posible_empate", "minuto": "50'", "texto": "x",
         "estado": "fallo", "acierto": False},
        {"tipo": "value_alert", "minuto": "60'", "texto": "x",
         "estado": "ambiguo", "acierto": None},
    ]
    C._actualizar_excel("2026-09-18", [p, p2], 10)
    wb = openpyxl.load_workbook(xlsx)
    hoja3 = wb["Acierto por tipo de alerta"]
    filas = {(r[1].value): [c.value for c in r] for r in hoja3.iter_rows(min_row=2)}
    assert "posible_empate" not in filas  # unificado bajo el nombre nuevo
    des = filas["posible_descuento"]
    assert des[2] == 1 and des[3] == 0 and des[5] == 1  # enviadas, aciertos, fallos
    assert des[4] == 0.0  # % con denominador aciertos+fallos
    val = filas["value_alert"]
    assert val[2] == 1 and val[6] == 1  # ambigua -> Sin determinar, no fallo
    assert val[4] is None  # sin denominador -> sin % espurio


def test_excel_agrupa_mismo_tipo_mismo_partido(tmp_path, monkeypatch):
    """4 alertas 'gana_fav' en el mismo partido cuentan COMO 1 fila."""
    import openpyxl
    xlsx = tmp_path / "est.xlsx"
    monkeypatch.setattr(C, "ARCHIVO_EXCEL", xlsx)
    p = partido_con_alertas()
    p["acierto"] = True
    p["favorito"] = "Sevilla"
    p["cuota_inicial"] = ""
    p["probabilidad_inicial"] = ""
    p["alertas_enviadas"] = [
        {"tipo": "posible_victoria_favorito", "minuto": m, "texto": "x",
         "estado": "fallo", "acierto": False}
        for m in ("30'", "45'", "60'", "75'")
    ]
    C._actualizar_excel("2026-09-19", [p], 10)
    wb = openpyxl.load_workbook(xlsx)
    hoja3 = wb["Acierto por tipo de alerta"]
    filas = {(r[1].value): [c.value for c in r] for r in hoja3.iter_rows(min_row=2)}
    assert "posible_victoria_favorito" in filas
    row = filas["posible_victoria_favorito"]
    # enviadas=1 (grupo), fallos=1, aciertos=0
    assert row[2] == 1 and row[5] == 1 and row[3] == 0


def test_excel_agrupa_mixed_usa_mayoria(tmp_path, monkeypatch):
    """3 aciertos + 1 fallo del mismo tipo/partido → 1 acierto (mayoría)."""
    import openpyxl
    xlsx = tmp_path / "est2.xlsx"
    monkeypatch.setattr(C, "ARCHIVO_EXCEL", xlsx)
    p = partido_con_alertas()
    p["acierto"] = True
    p["favorito"] = "Sevilla"
    p["cuota_inicial"] = ""
    p["probabilidad_inicial"] = ""
    p["alertas_enviadas"] = [
        {"tipo": "posible_victoria_favorito", "minuto": "30'", "texto": "x",
         "estado": "acierto", "acierto": True},
        {"tipo": "posible_victoria_favorito", "minuto": "45'", "texto": "x",
         "estado": "acierto", "acierto": True},
        {"tipo": "posible_victoria_favorito", "minuto": "60'", "texto": "x",
         "estado": "acierto", "acierto": True},
        {"tipo": "posible_victoria_favorito", "minuto": "75'", "texto": "x",
         "estado": "fallo", "acierto": False},
    ]
    C._actualizar_excel("2026-09-20", [p], 10)
    wb = openpyxl.load_workbook(xlsx)
    hoja3 = wb["Acierto por tipo de alerta"]
    filas = {(r[1].value): [c.value for c in r] for r in hoja3.iter_rows(min_row=2)}
    row = filas["posible_victoria_favorito"]
    # 1 grupo, mayoría aciertos (3 vs 1)
    assert row[2] == 1 and row[3] == 1 and row[5] == 0
